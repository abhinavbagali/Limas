from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Mapping

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[3]
DEFAULT_WORKBOOK = ROOT / "reports" / "eval_met.xlsx"
WORKSHEET_NAME = "eval_met"
HEADERS = (
    "Timestamp (UTC)",
    "Test Case",
    "Query",
    "Model",
    "Retrieved Passages",
    "Without Compressor - Prompt Tokens (Groq actual)",
    "With Compressor - Prompt Tokens (Groq actual)",
    "Prompt Tokens Saved (Groq actual)",
    "Prompt Token Reduction (%) (Groq actual)",
    "Sentences Retained",
    "Compression Latency (ms)",
    "Without Compressor - LLM Latency (ms)",
    "With Compressor - LLM Latency (ms)",
    "Without Compressor - Total Latency (ms)",
    "With Compressor - Total Latency (ms)",
    "Latency Saved (ms)",
    "Without Compressor - Completion Tokens (API)",
    "With Compressor - Completion Tokens (API)",
    "Without Compressor - Answer",
    "With Compressor - Answer",
    "Expected Answer",
    "Without Compressor - Reference Similarity",
    "With Compressor - Reference Similarity",
    "Legacy Original Context Word Count",
    "Legacy Compressed Context Word Count",
    "Legacy Estimated Context Reduction (%)",
    "Legacy Answer Similarity (old benchmark)",
    "Query Complexity Score",
    "Query Intent Type",
    "Query Entity Count",
    "Dynamic Threshold Used",
    "Dynamic Token Ratio Limit",
    "Preserved Coreferences",
    "Retained Sentence Indices",
    "Isolated Scoring Latency (ms)",
    "Pairwise Scoring Latency (ms)",
    "Coreference Analysis Latency (ms)",
)


def _migrate_headers(worksheet: Any) -> None:
    existing_headers = [cell.value for cell in worksheet[1]]
    if existing_headers[: len(HEADERS)] == list(HEADERS):
        return
    if not existing_headers or all(value is None for value in existing_headers):
        for column, header in enumerate(HEADERS, start=1):
            worksheet.cell(row=1, column=column, value=header)
        return

    old_rows = list(worksheet.iter_rows(min_row=2, values_only=True))
    old_index = {header: index for index, header in enumerate(existing_headers) if header}
    aliases = {
        "Timestamp (UTC)": ("Timestamp (UTC)",),
        "Test Case": ("Test Case",),
        "Query": ("Query",),
        "Model": ("Model",),
        "Retrieved Passages": ("Retrieved Passages",),
        "Without Compressor - Prompt Tokens (Groq actual)": (
            "Without Compressor - Prompt Tokens (Groq actual)",
            "Without Compressor - Prompt Tokens (API)",
        ),
        "With Compressor - Prompt Tokens (Groq actual)": (
            "With Compressor - Prompt Tokens (Groq actual)",
            "With Compressor - Prompt Tokens (API)",
        ),
        "Sentences Retained": ("Sentences Retained",),
        "Compression Latency (ms)": ("Compression Latency (ms)",),
        "Without Compressor - LLM Latency (ms)": ("Without Compressor - LLM Latency (ms)",),
        "With Compressor - LLM Latency (ms)": ("With Compressor - LLM Latency (ms)",),
        "Without Compressor - Total Latency (ms)": ("Without Compressor - Total Latency (ms)",),
        "With Compressor - Total Latency (ms)": ("With Compressor - Total Latency (ms)",),
        "Latency Saved (ms)": ("Latency Saved (ms)",),
        "Without Compressor - Completion Tokens (API)": ("Without Compressor - Completion Tokens (API)",),
        "With Compressor - Completion Tokens (API)": ("With Compressor - Completion Tokens (API)",),
        "Without Compressor - Answer": ("Without Compressor - Answer",),
        "With Compressor - Answer": ("With Compressor - Answer",),
        "Expected Answer": ("Expected Answer",),
        "Without Compressor - Reference Similarity": ("Without Compressor - Reference Similarity",),
        "With Compressor - Reference Similarity": ("With Compressor - Reference Similarity",),
        "Legacy Original Context Word Count": ("Legacy Original Context Word Count", "Original Context Tokens (estimated)", "Original Tokens"),
        "Legacy Compressed Context Word Count": ("Legacy Compressed Context Word Count", "Compressed Context Tokens (estimated)", "Compressed Tokens"),
        "Legacy Estimated Context Reduction (%)": ("Legacy Estimated Context Reduction (%)", "Context Token Reduction (%)", "Token Reduction (%)"),
        "Legacy Answer Similarity (old benchmark)": ("Legacy Answer Similarity (old benchmark)", "Answer Similarity"),
        "Query Complexity Score": ("Query Complexity Score",),
        "Query Intent Type": ("Query Intent Type",),
        "Query Entity Count": ("Query Entity Count",),
        "Dynamic Threshold Used": ("Dynamic Threshold Used",),
        "Dynamic Token Ratio Limit": ("Dynamic Token Ratio Limit",),
        "Preserved Coreferences": ("Preserved Coreferences",),
        "Retained Sentence Indices": ("Retained Sentence Indices",),
        "Isolated Scoring Latency (ms)": ("Isolated Scoring Latency (ms)",),
        "Pairwise Scoring Latency (ms)": ("Pairwise Scoring Latency (ms)",),
        "Coreference Analysis Latency (ms)": ("Coreference Analysis Latency (ms)",),
    }
    for row_number, old_row in enumerate(old_rows, start=2):
        preserved: dict[str, Any] = {}
        for new_header, old_headers in aliases.items():
            for old_header in old_headers:
                index = old_index.get(old_header)
                if index is not None and index < len(old_row):
                    preserved[new_header] = old_row[index]
                    break
        baseline_prompt_tokens = preserved.get(
            "Without Compressor - Prompt Tokens (Groq actual)"
        )
        compressed_prompt_tokens = preserved.get(
            "With Compressor - Prompt Tokens (Groq actual)"
        )
        if isinstance(baseline_prompt_tokens, (int, float)) and isinstance(
            compressed_prompt_tokens, (int, float)
        ):
            saved_prompt_tokens = baseline_prompt_tokens - compressed_prompt_tokens
            preserved["Prompt Tokens Saved (Groq actual)"] = saved_prompt_tokens
            if baseline_prompt_tokens > 0:
                preserved["Prompt Token Reduction (%) (Groq actual)"] = (
                    saved_prompt_tokens / baseline_prompt_tokens * 100.0
                )
        for column, header in enumerate(HEADERS, start=1):
            worksheet.cell(row=row_number, column=column, value=preserved.get(header))

    for column, header in enumerate(HEADERS, start=1):
        worksheet.cell(row=1, column=column, value=header)


def _format_worksheet(worksheet: Any) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(1, worksheet.max_row)}"
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[1].height = 38
    widths = (22, 22, 44, 25, 18, 34, 34, 29, 34, 17, 22, 30, 28, 31, 29, 18, 34, 32, 60, 60, 50, 34, 32, 32, 34, 36, 36, 24, 20, 18, 23, 25, 22, 28, 27, 27, 31)
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            if cell.column in (9, 22, 23, 26, 27):
                cell.number_format = "0.00"
            elif cell.column in (28, 31, 32):
                cell.number_format = "0.000"
            elif cell.column in range(11, 17):
                cell.number_format = "0.00"
            elif cell.column in range(35, 38):
                cell.number_format = "0.00"
            elif cell.column in range(19, 28) or cell.column == 34:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if worksheet.max_column >= 19:
            worksheet.row_dimensions[row[0].row].height = 48


def migrate_workbook(workbook_path: str | Path = DEFAULT_WORKBOOK) -> Path:
    """Create the comparison worksheet or migrate its headers without appending a run."""
    path = Path(workbook_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        workbook = load_workbook(path)
    else:
        workbook = Workbook()
        default_sheet = workbook.active
        if default_sheet is not None:
            workbook.remove(default_sheet)

    worksheet = workbook[WORKSHEET_NAME] if WORKSHEET_NAME in workbook.sheetnames else workbook.create_sheet(WORKSHEET_NAME)
    _migrate_headers(worksheet)
    _format_worksheet(worksheet)
    workbook.active = workbook.sheetnames.index(WORKSHEET_NAME)
    workbook.save(path)
    return path


def append_comparison_metrics(
    metrics: Mapping[str, Any],
    workbook_path: str | Path = DEFAULT_WORKBOOK,
) -> Path:
    """Append one paired baseline/compressed LLM evaluation row."""
    path = Path(workbook_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    if path.exists():
        workbook = load_workbook(path)
    else:
        workbook = Workbook()
        default_sheet = workbook.active
        if default_sheet is not None:
            workbook.remove(default_sheet)

    worksheet = workbook[WORKSHEET_NAME] if WORKSHEET_NAME in workbook.sheetnames else workbook.create_sheet(WORKSHEET_NAME)
    _migrate_headers(worksheet)
    row = {
        "Timestamp (UTC)": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        **metrics,
    }
    worksheet.append([row.get(header) for header in HEADERS])
    _format_worksheet(worksheet)
    workbook.active = workbook.sheetnames.index(WORKSHEET_NAME)
    workbook.save(path)
    return path
