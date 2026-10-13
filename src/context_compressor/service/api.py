from __future__ import annotations

import asyncio
import io
import logging
import os
import threading
import time
from contextlib import asynccontextmanager
from dataclasses import asdict
from pathlib import Path
from typing import List, Optional
from uuid import uuid4

from dotenv import load_dotenv
from fastapi.responses import FileResponse
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.staticfiles import StaticFiles
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from ..compression.compressor import ContextCompressor
from ..ingestion.chunker import TextChunk, chunk_pages
from ..ingestion.pdf_extractor import PageMetadata, PdfIngestionError, extract_pdf_pages
from ..retrieval.reranker import CrossEncoderReranker, MS_MARCO_MODEL
from ..retrieval.vector_store import RetrievedChunk, VectorStore
from ..evaluation.benchmark import _jaccard_similarity
from ..models.cross_encoder import AutoModelForSequenceClassification, AutoTokenizer, torch
from ..evaluation.comparison_metrics import (
    DEFAULT_WORKBOOK,
    WORKSHEET_NAME,
    append_comparison_metrics,
)
from .groq_client import (
    GroqConfigurationError,
    GroqRequestError,
    generate_answer,
    get_groq_model,
    estimate_prompt_tokens,
    fit_context_to_prompt_budget,
    MAX_COMPLETION_TOKENS,
)

logger = logging.getLogger(__name__)
PROJECT_ROOT = Path(__file__).resolve().parents[3]
load_dotenv(PROJECT_ROOT / ".env")

class CompressRequest(BaseModel):
    query: str = Field(..., description="User query to evaluate the retrieved context against.")
    documents: List[str] = Field(default_factory=list, description="Retrieved document chunks or passages.")
    doc_id: Optional[str] = None
    retrieval_top_k: int = Field(default=30, ge=1, le=200)
    rerank_top_k: int = Field(default=8, ge=1, le=50)
    token_aware_budget: bool = True
    threshold: Optional[float] = Field(default=None, ge=0.0, le=1.0)
    top_k: Optional[int] = Field(default=None, ge=1)
    max_tokens: Optional[int] = Field(default=None, ge=1)
    enable_coref: bool = True
    enable_multihop: bool = True
    adaptive_gating: bool = True


class SentenceAuditResponse(BaseModel):
    sentence_index: int
    text: str
    isolated_score: float
    pairwise_gain: float
    final_weight: float
    retained: bool
    atomic_group_id: Optional[int]
    retention_reason: Optional[str]
    discard_reason: Optional[str]
    doc_id: Optional[str] = None
    page_start: Optional[int] = None
    page_end: Optional[int] = None
    retrieval_rank: Optional[int] = None
    retrieval_score: Optional[float] = None
    rerank_score: Optional[float] = None


class CompressResponse(BaseModel):
    compressed_text: str
    selected_sentences: List[str]
    scores: List[float]
    sentence_audit: List[SentenceAuditResponse]
    estimated_original_word_count: int
    estimated_compressed_word_count: int
    estimated_word_reduction_percent: float
    compression_ratio: float
    query_complexity_score: float
    query_intent_type: str
    query_entity_count: int
    complexity_signals: List[str]
    dynamic_threshold_used: float
    dynamic_token_ratio_limit: Optional[float]
    preserved_coreferences_count: int
    retained_sentence_indices: List[int]
    isolated_scoring_latency_ms: float
    pairwise_scoring_latency_ms: float
    coreference_analysis_latency_ms: float
    retrieved_chunks: List["RetrievalItemResponse"] = Field(default_factory=list)


class CompareRequest(BaseModel):
    query: str = Field(..., min_length=1)
    documents: List[str] = Field(default_factory=list)
    doc_id: Optional[str] = None
    retrieval_top_k: int = Field(default=30, ge=1, le=200)
    rerank_top_k: int = Field(default=8, ge=1, le=50)
    tpm_limit: int = Field(default=8000, ge=1024, le=100000)
    token_aware_budget: bool = True
    test_case: Optional[str] = None
    expected_answer: Optional[str] = None
    threshold: Optional[float] = Field(default=None, ge=0.0, le=1.0)
    top_k: Optional[int] = Field(default=None, ge=1)
    max_tokens: Optional[int] = Field(default=None, ge=1)
    enable_coref: bool = True
    enable_multihop: bool = True
    adaptive_gating: bool = True


class CompareResponse(BaseModel):
    model: str
    test_case: Optional[str]
    compressed_text: str
    selected_sentences: List[str]
    scores: List[float]
    sentence_audit: List[SentenceAuditResponse]
    estimated_original_word_count: int
    estimated_compressed_word_count: int
    estimated_word_reduction_percent: float
    compression_ratio: float
    query_complexity_score: float
    query_intent_type: str
    query_entity_count: int
    complexity_signals: List[str]
    dynamic_threshold_used: float
    dynamic_token_ratio_limit: Optional[float]
    preserved_coreferences_count: int
    retained_sentence_indices: List[int]
    isolated_scoring_latency_ms: float
    pairwise_scoring_latency_ms: float
    coreference_analysis_latency_ms: float
    compression_latency_ms: float
    without_compressor_answer: str
    with_compressor_answer: str
    without_compressor_llm_latency_ms: float
    with_compressor_llm_latency_ms: float
    without_compressor_prompt_tokens: Optional[int]
    with_compressor_prompt_tokens: Optional[int]
    actual_prompt_tokens_saved: Optional[int]
    actual_prompt_token_reduction_percent: Optional[float]
    without_compressor_completion_tokens: Optional[int]
    with_compressor_completion_tokens: Optional[int]
    without_compressor_answer_similarity: Optional[float]
    with_compressor_answer_similarity: Optional[float]
    metrics_saved: bool
    metrics_save_error: Optional[str]
    retrieved_chunks: List["RetrievalItemResponse"] = Field(default_factory=list)
    preflight_prompt_token_budget: int
    baseline_preflight_prompt_tokens: int
    compressed_preflight_prompt_tokens: int
    baseline_context_truncated: bool


class UploadedDocumentResponse(BaseModel):
    filename: str
    passages: list[str]
    page_count: int
    doc_id: str
    pages: list[dict[str, object]]
    indexed: bool = False


class IndexRequest(BaseModel):
    doc_id: str
    target_tokens: int = Field(default=500, ge=100, le=1000)
    overlap_tokens: int = Field(default=80, ge=0, le=200)


class RetrieveRequest(BaseModel):
    doc_id: str
    query: str = Field(..., min_length=1)
    top_k: int = Field(default=30, ge=1, le=200)
    rerank_top_k: int = Field(default=8, ge=1, le=50)
    reranker: str = MS_MARCO_MODEL


class RetrievalItemResponse(BaseModel):
    chunk_id: str
    doc_id: str
    text: str
    page_start: int
    page_end: int
    section_title: Optional[str]
    token_count: int
    retrieval_score: float
    retrieval_rank: int
    rerank_score: Optional[float] = None
    rerank_rank: Optional[int] = None


def _numeric_values(rows: list[dict[str, object]], key: str) -> list[float]:
    return [
        float(value)
        for row in rows
        if isinstance((value := row.get(key)), (int, float)) and not isinstance(value, bool)
    ]


def _average(values: list[float]) -> Optional[float]:
    return sum(values) / len(values) if values else None


def _load_analytics_rows() -> list[dict[str, object]]:
    if not DEFAULT_WORKBOOK.exists():
        return []
    workbook = load_workbook(DEFAULT_WORKBOOK, read_only=True, data_only=True)
    try:
        if WORKSHEET_NAME not in workbook.sheetnames:
            return []
        worksheet = workbook[WORKSHEET_NAME]
        records = worksheet.iter_rows(values_only=True)
        headers = next(records, ())
        return [
            {str(header): value for header, value in zip(headers, values) if header}
            for values in records
        ]
    finally:
        workbook.close()


def _load_paired_analytics_rows() -> list[dict[str, object]]:
    return [
        row
        for row in _load_analytics_rows()
        if isinstance(row.get("Without Compressor - Total Latency (ms)"), (int, float))
    ]


@asynccontextmanager
async def lifespan(app: FastAPI):
    """Load retrieval models once when the API process starts."""
    vector_store._load_model()
    _get_reranker(MS_MARCO_MODEL)
    yield


app = FastAPI(title="Limas — Subcortical Context Compressor", version="0.1.0", lifespan=lifespan)
compressor = ContextCompressor()
vector_store = VectorStore()
rerankers: dict[str, CrossEncoderReranker] = {}
document_pages: dict[str, list[PageMetadata]] = {}
document_chunks: dict[str, list[TextChunk]] = {}
document_chunk_config: dict[str, tuple[int, int]] = {}
UI_PATH = Path(__file__).parent / "static" / "index.html"
UI_DIST_DIR = PROJECT_ROOT / "ui" / "dist"
UI_CLIENT_DIR = UI_DIST_DIR / "client"
UI_STATIC_DIR = (
    UI_CLIENT_DIR
    if (UI_CLIENT_DIR / "index.html").is_file()
    else UI_DIST_DIR
)


def _get_reranker(model_name: str) -> CrossEncoderReranker:
    if model_name not in rerankers:
        rerankers.clear()
        rerankers[model_name] = CrossEncoderReranker(model_name)
    return rerankers[model_name]


def _resolve_context(doc_id: str, query: str, retrieval_top_k: int, rerank_top_k: int) -> tuple[list[str], list[RetrievedChunk], list[object]]:
    candidates = vector_store.retrieve(doc_id, query, retrieval_top_k)
    reranked = _get_reranker(MS_MARCO_MODEL).rerank(query, candidates, rerank_top_k)
    selected: list[RetrievedChunk] = []
    by_chunk = {item.chunk.chunk_id: item for item in candidates}
    for item in reranked:
        source = by_chunk[item.chunk.chunk_id]
        selected.append(RetrievedChunk(item.chunk, source.score, item.retrieval_rank))
    return [item.chunk.text for item in reranked], selected, reranked


def _audit_provenance(audit: list[SentenceAuditResponse], ranked_chunks: list[object]) -> None:
    from ..compression.sentence_splitter import split_document_into_sentences

    provenance: list[tuple[str, int, int, int, float, float]] = []
    for item in ranked_chunks:
        chunk = item.chunk
        for _ in split_document_into_sentences(chunk.text):
            provenance.append((chunk.doc_id, chunk.page_start, chunk.page_end, item.retrieval_rank, item.retrieval_score, item.rerank_score))
    for entry in audit:
        if 0 <= entry.sentence_index < len(provenance):
            entry.doc_id, entry.page_start, entry.page_end, entry.retrieval_rank, entry.retrieval_score, entry.rerank_score = provenance[entry.sentence_index]


def _retrieval_response(ranked_chunks: list[object]) -> list[RetrievalItemResponse]:
    return [
        RetrievalItemResponse(
            chunk_id=item.chunk.chunk_id, doc_id=item.chunk.doc_id, text=item.chunk.text,
            page_start=item.chunk.page_start, page_end=item.chunk.page_end,
            section_title=item.chunk.section_title, token_count=item.chunk.token_count,
            retrieval_score=item.retrieval_score, retrieval_rank=item.retrieval_rank,
            rerank_score=item.rerank_score, rerank_rank=item.rerank_rank,
        ) for item in ranked_chunks
    ]


@app.get("/", include_in_schema=False)
async def pipeline_ui() -> FileResponse:
    react_index = UI_STATIC_DIR / "index.html"
    return FileResponse(react_index if react_index.is_file() else UI_PATH)


@app.get("/health")
async def health() -> dict[str, str]:
    return {"status": "ok", "service": "limas"}


@app.get("/config")
async def runtime_config() -> dict[str, str | bool]:
    reranker = rerankers.get(MS_MARCO_MODEL)
    embedding_ready = vector_store._model is not None and vector_store._model is not False
    return {
        "groq_configured": bool(os.getenv("GROQ_API_KEY")),
        "model": get_groq_model(),
        "transformers_available": torch is not None and AutoTokenizer is not None and AutoModelForSequenceClassification is not None,
        "embedding_model_loaded": embedding_ready,
        "embedding_model": vector_store.model_name,
        "reranker_model_loaded": bool(reranker and not reranker.scorer.fallback_mode),
        "reranker_model": MS_MARCO_MODEL,
        "compression_scorer_loaded": not compressor.scorer.fallback_mode,
        "compression_scorer_model": compressor.scorer.model_name,
    }


@app.post("/upload_pdf", response_model=UploadedDocumentResponse)
def upload_document(file: UploadFile = File(...)) -> UploadedDocumentResponse:
    filename = Path(file.filename or "").name
    extension = Path(filename).suffix.lower()
    if extension not in {".pdf", ".txt"}:
        raise HTTPException(status_code=415, detail="Upload a PDF or plain-text (.txt) file.")

    content = file.file.read(25 * 1024 * 1024 + 1)
    if len(content) > 25 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="Uploaded files must be 25 MB or smaller.")
    if not content:
        raise HTTPException(status_code=422, detail="The uploaded file is empty.")

    if extension == ".txt":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise HTTPException(status_code=422, detail="Text files must use UTF-8 encoding.") from exc
        normalized_text = text.replace("\r\n", "\n").replace("\r", "\n")
        passages = [part.strip() for part in normalized_text.split("\n\n") if part.strip()]
        page_count = 1
        doc_id = str(uuid4())
        pages = [PageMetadata(doc_id, 1, normalized_text.strip(), "success" if normalized_text.strip() else "empty")]
    else:
        try:
            pages = extract_pdf_pages(content, reader_factory=PdfReader)
            doc_id = pages[0].doc_id
            passages = [page.text for page in pages if page.text.strip()]
        except PdfIngestionError as exc:
            logger.info("Could not parse uploaded PDF %s: %s", filename, exc)
            raise HTTPException(status_code=422, detail=str(exc)) from exc
        page_count = len(pages)

    if not passages:
        raise HTTPException(
            status_code=422,
            detail="No extractable text was found. Scanned PDFs require OCR before upload.",
        )
    document_pages[doc_id] = pages
    document_chunks.pop(doc_id, None)
    document_chunk_config.pop(doc_id, None)
    return UploadedDocumentResponse(
        filename=filename,
        passages=passages,
        page_count=page_count,
        doc_id=doc_id,
        pages=[asdict(page) for page in pages],
    )


@app.post("/index")
def index_document(request: IndexRequest) -> dict[str, object]:
    pages = document_pages.get(request.doc_id)
    if pages is None:
        raise HTTPException(status_code=404, detail="Document was not found in this API process. Upload it again.")
    process = None
    stop_sampling = threading.Event()
    peak_rss_bytes = 0
    try:
        import psutil

        process = psutil.Process()
        peak_rss_bytes = process.memory_info().rss

        def sample_rss() -> None:
            nonlocal peak_rss_bytes
            while not stop_sampling.wait(0.05):
                try:
                    peak_rss_bytes = max(peak_rss_bytes, process.memory_info().rss)
                except Exception:
                    logger.exception("Could not sample process RSS during indexing")
                    return

        sampler = threading.Thread(target=sample_rss, name="index-rss-sampler", daemon=True)
        sampler.start()
    except ImportError:
        sampler = None
        logger.warning("psutil is unavailable; skipping indexing memory profile")
    rss_before = peak_rss_bytes
    try:
        chunks = chunk_pages(pages, request.target_tokens, request.overlap_tokens)
        force_reindex = document_chunk_config.get(request.doc_id) != (
            request.target_tokens,
            request.overlap_tokens,
        )
        index = vector_store.index(request.doc_id, chunks, force=force_reindex)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    finally:
        stop_sampling.set()
        if sampler is not None:
            sampler.join(timeout=1)
        if process is not None:
            peak_rss_bytes = max(peak_rss_bytes, process.memory_info().rss)
    document_chunks[request.doc_id] = chunks
    document_chunk_config[request.doc_id] = (request.target_tokens, request.overlap_tokens)
    memory_profile = None
    if process is not None:
        memory_profile = {
            "start_rss_mb": round(rss_before / (1024 * 1024), 2),
            "peak_rss_mb": round(peak_rss_bytes / (1024 * 1024), 2),
            "end_rss_mb": round(process.memory_info().rss / (1024 * 1024), 2),
        }
        logger.info("PDF indexing RSS profile doc_id=%s profile=%s", request.doc_id, memory_profile)
    return {
        "doc_id": request.doc_id,
        "page_count": len(pages),
        "chunk_count": len(chunks),
        "target_tokens": request.target_tokens,
        "overlap_tokens": request.overlap_tokens,
        "embedding_backend": index.backend,
        "embedding_model": vector_store.model_name if index.backend == "sentence-transformers" else None,
        "retrieval_method": "dense + BM25 with exact technical-term matching",
        "memory_profile": memory_profile,
    }


@app.post("/retrieve", response_model=list[RetrievalItemResponse])
def retrieve_document(request: RetrieveRequest) -> list[RetrievalItemResponse]:
    if not vector_store.has_index(request.doc_id):
        raise HTTPException(status_code=409, detail="Document is not indexed. Call POST /index first.")
    candidates = vector_store.retrieve(request.doc_id, request.query, request.top_k)
    try:
        reranked = _get_reranker(request.reranker).rerank(request.query, candidates, request.rerank_top_k)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    retrieval_by_id = {item.chunk.chunk_id: item for item in candidates}
    return [
        RetrievalItemResponse(
            chunk_id=item.chunk.chunk_id,
            doc_id=item.chunk.doc_id,
            text=item.chunk.text,
            page_start=item.chunk.page_start,
            page_end=item.chunk.page_end,
            section_title=item.chunk.section_title,
            token_count=item.chunk.token_count,
            retrieval_score=retrieval_by_id[item.chunk.chunk_id].score,
            retrieval_rank=retrieval_by_id[item.chunk.chunk_id].rank,
            rerank_score=item.rerank_score,
            rerank_rank=item.rerank_rank,
        )
        for item in reranked
    ]


@app.get("/analytics")
def analytics() -> dict[str, object]:
    rows = _load_paired_analytics_rows()
    prompt_reductions = _numeric_values(rows, "Prompt Token Reduction (%) (Groq actual)")
    latency_savings = _numeric_values(rows, "Latency Saved (ms)")
    quality_scores = _numeric_values(rows, "With Compressor - Reference Similarity")
    coreference_counts = _numeric_values(rows, "Preserved Coreferences")
    intents: dict[str, int] = {}
    for row in rows:
        intent = row.get("Query Intent Type")
        if isinstance(intent, str) and intent.strip():
            label = intent.strip().replace("_", "-").title()
            intents[label] = intents.get(label, 0) + 1
    intent_distribution = [
        {"name": name, "count": count}
        for name, count in sorted(intents.items(), key=lambda item: item[0])
    ]

    history = []
    for index, row in enumerate(rows[-105:], start=max(1, len(rows) - 104)):
        saved = row.get("Prompt Token Reduction (%) (Groq actual)")
        history.append({
            "run_number": index,
            "timestamp": row.get("Timestamp (UTC)"),
            "saved_percent": float(saved) if isinstance(saved, (int, float)) else None,
            "coref_edges": int(row["Preserved Coreferences"]) if isinstance(row.get("Preserved Coreferences"), (int, float)) else 0,
            "latency_saved_ms": float(row["Latency Saved (ms)"]) if isinstance(row.get("Latency Saved (ms)"), (int, float)) else None,
        })

    chart_rows = []
    for index, row in enumerate(rows[-10:], start=max(1, len(rows) - 9)):
        chart_rows.append({
            "run_number": index,
            "test_case": row.get("Test Case") or f"Run {index}",
            "baseline_prompt_tokens": row.get("Without Compressor - Prompt Tokens (Groq actual)"),
            "compressed_prompt_tokens": row.get("With Compressor - Prompt Tokens (Groq actual)"),
            "latency_saved_ms": row.get("Latency Saved (ms)"),
        })

    return {
        "runs_total": len(rows),
        "avg_prompt_token_reduction_percent": _average(prompt_reductions),
        "avg_latency_saved_ms": _average(latency_savings),
        "avg_quality_retention_percent": (
            _average(quality_scores) * 100 if quality_scores else None
        ),
        "coreferences_preserved": int(sum(coreference_counts)),
        "activity": history,
        "performance": chart_rows,
        "intent_distribution": intent_distribution,
        "dynamic_threshold_range": {
            "min": min(_numeric_values(rows, "Dynamic Threshold Used"), default=None),
            "max": max(_numeric_values(rows, "Dynamic Threshold Used"), default=None),
        },
    }


@app.post("/compress", response_model=CompressResponse)
async def compress(request: CompressRequest) -> CompressResponse:
    ranked_chunks: list[object] = []
    documents = request.documents
    if request.doc_id:
        if not vector_store.has_index(request.doc_id):
            raise HTTPException(status_code=409, detail="Document is not indexed. Call POST /index first.")
        documents, _, ranked_chunks = _resolve_context(request.doc_id, request.query, request.retrieval_top_k, request.rerank_top_k)
    elif not documents:
        raise HTTPException(status_code=422, detail="Provide documents or a doc_id.")
    token_counter = None
    if request.doc_id and request.token_aware_budget:
        token_counter = lambda text: max(
            0,
            estimate_prompt_tokens(request.query, text)
            - estimate_prompt_tokens(request.query, ""),
        )
    result = compressor.compress(
        query=request.query,
        documents=documents,
        threshold=request.threshold,
        top_k=request.top_k,
        max_tokens=request.max_tokens,
        enable_coref=request.enable_coref,
        enable_multihop=request.enable_multihop,
        adaptive_gating=request.adaptive_gating,
        token_counter=token_counter,
    )
    sentence_audit = [SentenceAuditResponse(**asdict(item)) for item in result.sentence_audit]
    if ranked_chunks:
        _audit_provenance(sentence_audit, ranked_chunks)
    return CompressResponse(
        compressed_text=result.compressed_text,
        selected_sentences=result.selected_sentences,
        scores=result.scores,
        sentence_audit=sentence_audit,
        estimated_original_word_count=result.estimated_original_word_count,
        estimated_compressed_word_count=result.estimated_compressed_word_count,
        estimated_word_reduction_percent=result.estimated_word_reduction_percent,
        compression_ratio=result.compression_ratio,
        query_complexity_score=result.query_complexity_score,
        query_intent_type=result.query_intent_type,
        query_entity_count=result.query_entity_count,
        complexity_signals=result.complexity_signals,
        dynamic_threshold_used=result.dynamic_threshold_used,
        dynamic_token_ratio_limit=result.dynamic_token_ratio_limit,
        preserved_coreferences_count=result.preserved_coreferences_count,
        retained_sentence_indices=result.retained_sentence_indices,
        isolated_scoring_latency_ms=result.isolated_scoring_latency_ms,
        pairwise_scoring_latency_ms=result.pairwise_scoring_latency_ms,
        coreference_analysis_latency_ms=result.coreference_analysis_latency_ms,
        retrieved_chunks=_retrieval_response(ranked_chunks),
    )


@app.post("/compare", response_model=CompareResponse)
async def compare(request: CompareRequest) -> CompareResponse:
    if not request.query.strip():
        raise HTTPException(status_code=422, detail="Query must not be blank.")
    ranked_chunks: list[object] = []
    documents = request.documents
    if request.doc_id:
        if not vector_store.has_index(request.doc_id):
            raise HTTPException(status_code=409, detail="Document is not indexed. Call POST /index first.")
        documents, _, ranked_chunks = _resolve_context(request.doc_id, request.query, request.retrieval_top_k, request.rerank_top_k)
    elif not documents:
        raise HTTPException(status_code=422, detail="Provide documents or a doc_id.")
    if any(not document.strip() for document in documents):
        raise HTTPException(status_code=422, detail="Passages must not contain blank entries.")

    # Split the available TPM budget across both concurrent calls, reserving
    # each call's configured completion allowance before fitting prompt text.
    prompt_budget = max(1, (request.tpm_limit - 2 * MAX_COMPLETION_TOKENS) // 2)
    if estimate_prompt_tokens(request.query, "") > prompt_budget:
        raise HTTPException(
            status_code=422,
            detail="The query and fixed prompt exceed the paired TPM budget; shorten the query or raise tpm_limit.",
        )
    documents, baseline_context_truncated = fit_context_to_prompt_budget(
        request.query, documents, prompt_budget
    )
    if baseline_context_truncated:
        logger.warning(
            "Truncated baseline context to estimated prompt budget (%s tokens) for test case %r",
            prompt_budget,
            request.test_case,
        )

    compression_start = time.perf_counter()
    token_counter = None
    if request.token_aware_budget:
        token_counter = lambda text: max(
            0,
            estimate_prompt_tokens(request.query, text)
            - estimate_prompt_tokens(request.query, ""),
        )
    compressed = compressor.compress(
        query=request.query,
        documents=documents,
        threshold=request.threshold,
        top_k=request.top_k,
        max_tokens=min(
            request.max_tokens or compressor.config.max_context_tokens,
            max(1, prompt_budget - estimate_prompt_tokens(request.query, "")),
        ),
        enable_coref=request.enable_coref,
        enable_multihop=request.enable_multihop,
        adaptive_gating=request.adaptive_gating,
        token_counter=token_counter,
    )
    compression_latency_ms = (time.perf_counter() - compression_start) * 1000.0
    original_context = "\n\n".join(documents)
    baseline_preflight_tokens = estimate_prompt_tokens(request.query, original_context)
    compressed_preflight_tokens = estimate_prompt_tokens(request.query, compressed.compressed_text)
    if baseline_preflight_tokens > prompt_budget or compressed_preflight_tokens > prompt_budget:
        raise HTTPException(
            status_code=422,
            detail=(
                "The paired contexts exceed the estimated TPM allowance after compression. "
                "Reduce the requested context budget or raise tpm_limit."
            ),
        )

    generation_results = await asyncio.gather(
        generate_answer(request.query, original_context),
        generate_answer(request.query, compressed.compressed_text),
        return_exceptions=True,
    )
    baseline_result, compressed_result = generation_results
    configuration_errors: list[str] = []
    request_errors: list[str] = []
    for branch, result in (
        ("baseline", baseline_result),
        ("compressed", compressed_result),
    ):
        if isinstance(result, GroqConfigurationError):
            configuration_errors.append(f"{branch}: {result}")
        elif isinstance(result, GroqRequestError):
            request_errors.append(f"{branch}: {result}")
        elif isinstance(result, BaseException):
            raise result
    if configuration_errors:
        raise HTTPException(
            status_code=503,
            detail="Groq configuration failed: " + "; ".join(configuration_errors),
        )
    if request_errors:
        raise HTTPException(
            status_code=502,
            detail="Groq generation failed: " + "; ".join(request_errors),
        )

    baseline = baseline_result
    compressed_answer = compressed_result

    baseline_similarity = (
        _jaccard_similarity(request.expected_answer, baseline.answer)
        if request.expected_answer and request.expected_answer.strip()
        else None
    )
    compressed_similarity = (
        _jaccard_similarity(request.expected_answer, compressed_answer.answer)
        if request.expected_answer and request.expected_answer.strip()
        else None
    )
    actual_prompt_tokens_saved = None
    actual_prompt_token_reduction_percent = None
    if baseline.prompt_tokens is not None and compressed_answer.prompt_tokens is not None:
        actual_prompt_tokens_saved = baseline.prompt_tokens - compressed_answer.prompt_tokens
        if baseline.prompt_tokens > 0:
            actual_prompt_token_reduction_percent = (
                actual_prompt_tokens_saved / baseline.prompt_tokens * 100.0
            )

    baseline_total_ms = baseline.latency_ms
    compressed_total_ms = compression_latency_ms + compressed_answer.latency_ms

    metric_values = {
        "Test Case": request.test_case or "",
        "Query": request.query,
        "Model": get_groq_model(),
        "Retrieved Passages": len(documents),
        "Sentences Retained": len(compressed.selected_sentences),
        "Query Complexity Score": compressed.query_complexity_score,
        "Query Intent Type": compressed.query_intent_type,
        "Dynamic Threshold Used": compressed.dynamic_threshold_used,
        "Dynamic Token Ratio Limit": compressed.dynamic_token_ratio_limit,
        "Preserved Coreferences": compressed.preserved_coreferences_count,
        "Retained Sentence Indices": ",".join(
            str(index) for index in compressed.retained_sentence_indices
        ),
        "Isolated Scoring Latency (ms)": compressed.isolated_scoring_latency_ms,
        "Pairwise Scoring Latency (ms)": compressed.pairwise_scoring_latency_ms,
        "Coreference Analysis Latency (ms)": compressed.coreference_analysis_latency_ms,
        "Compression Latency (ms)": compression_latency_ms,
        "Without Compressor - LLM Latency (ms)": baseline.latency_ms,
        "With Compressor - LLM Latency (ms)": compressed_answer.latency_ms,
        "Without Compressor - Total Latency (ms)": baseline_total_ms,
        "With Compressor - Total Latency (ms)": compressed_total_ms,
        "Latency Saved (ms)": baseline_total_ms - compressed_total_ms,
        "Without Compressor - Prompt Tokens (Groq actual)": baseline.prompt_tokens,
        "With Compressor - Prompt Tokens (Groq actual)": compressed_answer.prompt_tokens,
        "Prompt Tokens Saved (Groq actual)": actual_prompt_tokens_saved,
        "Prompt Token Reduction (%) (Groq actual)": actual_prompt_token_reduction_percent,
        "Without Compressor - Completion Tokens (API)": baseline.completion_tokens,
        "With Compressor - Completion Tokens (API)": compressed_answer.completion_tokens,
        "Without Compressor - Answer": baseline.answer,
        "With Compressor - Answer": compressed_answer.answer,
        "Expected Answer": request.expected_answer or "",
        "Without Compressor - Reference Similarity": baseline_similarity,
        "With Compressor - Reference Similarity": compressed_similarity,
    }

    metrics_saved = True
    metrics_save_error = None
    try:
        append_comparison_metrics(metric_values)
    except (OSError, PermissionError) as exc:
        logger.exception("Could not save comparison metrics workbook")
        metrics_saved = False
        metrics_save_error = (
            "LLM comparison succeeded, but metrics could not be saved. "
            "Close reports/eval_met.xlsx if it is open and try the comparison again."
        )

    sentence_audit = [SentenceAuditResponse(**asdict(item)) for item in compressed.sentence_audit]
    if ranked_chunks:
        _audit_provenance(sentence_audit, ranked_chunks)
    return CompareResponse(
        model=get_groq_model(),
        test_case=request.test_case,
        compressed_text=compressed.compressed_text,
        selected_sentences=compressed.selected_sentences,
        scores=compressed.scores,
        sentence_audit=sentence_audit,
        estimated_original_word_count=compressed.estimated_original_word_count,
        estimated_compressed_word_count=compressed.estimated_compressed_word_count,
        estimated_word_reduction_percent=compressed.estimated_word_reduction_percent,
        compression_ratio=compressed.compression_ratio,
        query_complexity_score=compressed.query_complexity_score,
        query_intent_type=compressed.query_intent_type,
        query_entity_count=compressed.query_entity_count,
        complexity_signals=compressed.complexity_signals,
        dynamic_threshold_used=compressed.dynamic_threshold_used,
        dynamic_token_ratio_limit=compressed.dynamic_token_ratio_limit,
        preserved_coreferences_count=compressed.preserved_coreferences_count,
        retained_sentence_indices=compressed.retained_sentence_indices,
        isolated_scoring_latency_ms=compressed.isolated_scoring_latency_ms,
        pairwise_scoring_latency_ms=compressed.pairwise_scoring_latency_ms,
        coreference_analysis_latency_ms=compressed.coreference_analysis_latency_ms,
        compression_latency_ms=compression_latency_ms,
        without_compressor_answer=baseline.answer,
        with_compressor_answer=compressed_answer.answer,
        without_compressor_llm_latency_ms=baseline.latency_ms,
        with_compressor_llm_latency_ms=compressed_answer.latency_ms,
        without_compressor_prompt_tokens=baseline.prompt_tokens,
        with_compressor_prompt_tokens=compressed_answer.prompt_tokens,
        actual_prompt_tokens_saved=actual_prompt_tokens_saved,
        actual_prompt_token_reduction_percent=actual_prompt_token_reduction_percent,
        without_compressor_completion_tokens=baseline.completion_tokens,
        with_compressor_completion_tokens=compressed_answer.completion_tokens,
        without_compressor_answer_similarity=baseline_similarity,
        with_compressor_answer_similarity=compressed_similarity,
        metrics_saved=metrics_saved,
        metrics_save_error=metrics_save_error,
        retrieved_chunks=_retrieval_response(ranked_chunks),
        preflight_prompt_token_budget=prompt_budget,
        baseline_preflight_prompt_tokens=baseline_preflight_tokens,
        compressed_preflight_prompt_tokens=compressed_preflight_tokens,
        baseline_context_truncated=baseline_context_truncated,
    )


# Keep this catch-all mount after the API routes so paths such as /compress,
# /compare, /index, and /retrieve continue to resolve to their API handlers.
# The root route above serves index.html; this mount serves its built assets.
if UI_STATIC_DIR.is_dir() and (UI_STATIC_DIR / "index.html").is_file():
    app.mount(
        "/",
        StaticFiles(directory=str(UI_STATIC_DIR), html=True),
        name="ui_dist",
    )
