"""Run a local four-way retrieval/compression benchmark from PDF/query cases.

Input JSON is a list of objects with `pdf`, `query`, and optional `answer` and `case_id`.
This script does not call Groq; its quality values are lexical context/reference proxies.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import time
from pathlib import Path

from openpyxl import load_workbook, Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from context_compressor.compression.compressor import ContextCompressor
from context_compressor.ingestion.chunker import chunk_pages
from context_compressor.ingestion.pdf_extractor import extract_pdf_pages
from context_compressor.retrieval.reranker import CrossEncoderReranker
from context_compressor.retrieval.vector_store import VectorStore
from context_compressor.service.groq_client import estimate_prompt_tokens


def words(text: str) -> set[str]:
    return set(re.findall(r"\w+", text.lower()))


def jaccard(left: str, right: str) -> float | None:
    a, b = words(left), words(right)
    return len(a & b) / len(a | b) if a and b else None


def tier_for(page_count: int) -> str:
    if 4 <= page_count <= 10:
        return "short"
    if 20 <= page_count <= 30:
        return "medium"
    if page_count >= 50:
        return "long"
    return "out_of_scope"


def run_case(case: dict[str, object], compressor: ContextCompressor, vector_store: VectorStore, reranker: CrossEncoderReranker) -> list[dict[str, object]]:
    pdf_path = (ROOT / str(case["pdf"])).resolve()
    pages = extract_pdf_pages(pdf_path.read_bytes())
    query = str(case["query"])
    answer = str(case.get("answer", ""))
    chunks = chunk_pages(pages)
    doc_id = pages[0].doc_id
    vector_store.index(doc_id, chunks)
    start = time.perf_counter()
    candidates = vector_store.retrieve(doc_id, query, top_k=30)
    ranked = reranker.rerank(query, candidates, top_k=8)
    retrieval_ms = (time.perf_counter() - start) * 1000
    retrieved = [item.chunk.text for item in ranked]
    full_pages = [page.text for page in pages if page.text]
    outputs = []

    baseline_tokens = estimate_prompt_tokens(query, "\n\n".join(full_pages))

    def record(
        name: str,
        context: str,
        elapsed_ms: float,
        recall: float | None = None,
        retrieval_stage_ms: float = 0.0,
        compression_stage_ms: float = 0.0,
    ) -> None:
        context_tokens = estimate_prompt_tokens(query, context)
        outputs.append({
            "case_id": case.get("case_id", pdf_path.stem),
            "tier": tier_for(len(pages)),
            "page_count": len(pages),
            "configuration": name,
            "recall_at_k": recall,
            "quality_jaccard_proxy": jaccard(context, answer) if answer else None,
            "prompt_token_reduction_percent": (baseline_tokens - context_tokens) / baseline_tokens * 100 if baseline_tokens else None,
            "context_characters": len(context),
            "latency_ms": elapsed_ms,
            "retrieval_latency_ms": retrieval_stage_ms,
            "compression_latency_ms": compression_stage_ms,
        })

    direct_start = time.perf_counter()
    direct = compressor.compress(query, full_pages).compressed_text
    direct_ms = (time.perf_counter() - direct_start) * 1000
    record("A_legacy_limas_direct", direct, direct_ms, compression_stage_ms=direct_ms)

    reference_terms = words(answer)
    relevant_total = sum(bool(reference_terms & words(chunk.text)) for chunk in chunks)
    hits = sum(bool(reference_terms & words(item.chunk.text)) for item in ranked)
    recall = hits / relevant_total if relevant_total else None
    raw_retrieved = "\n\n".join(retrieved)
    record("B_retrieval_only", raw_retrieved, retrieval_ms, recall, retrieval_stage_ms=retrieval_ms)

    enhanced_start = time.perf_counter()
    enhanced = compressor.compress(query, retrieved).compressed_text
    enhanced_ms = (time.perf_counter() - enhanced_start) * 1000
    record("C_retrieval_plus_limas", enhanced, retrieval_ms + enhanced_ms, recall, retrieval_stage_ms=retrieval_ms, compression_stage_ms=enhanced_ms)

    alternative_start = time.perf_counter()
    # A transparent extractive baseline: highest reranked chunks, capped at four.
    alternative = "\n\n".join(retrieved[:4])
    alternative_ms = (time.perf_counter() - alternative_start) * 1000
    record("D_retrieval_plus_top4_baseline", alternative, retrieval_ms + alternative_ms, recall, retrieval_stage_ms=retrieval_ms, compression_stage_ms=alternative_ms)
    return outputs


def write_rows(rows: list[dict[str, object]], path: Path) -> None:
    if path.exists():
        workbook = load_workbook(path)
    else:
        workbook = Workbook()
    if "retrieval_benchmark" in workbook.sheetnames:
        sheet = workbook["retrieval_benchmark"]
    else:
        sheet = workbook.create_sheet("retrieval_benchmark")
    if sheet.max_row == 1 and sheet.cell(1, 1).value is None:
        sheet.append(list(rows[0].keys()) if rows else ["case_id"])
    headers = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]
    for row in rows:
        sheet.append([row.get(header) for header in headers])
    workbook.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cases", required=True, type=Path, help="JSON list with pdf, query, and optional answer")
    parser.add_argument("--workbook", type=Path, default=ROOT / "reports" / "eval_met.xlsx")
    args = parser.parse_args()
    cases = json.loads(args.cases.read_text(encoding="utf-8"))
    if not isinstance(cases, list) or not cases:
        parser.error("--cases must contain a non-empty JSON list.")
    compressor, vector_store, reranker = ContextCompressor(), VectorStore(), CrossEncoderReranker()
    results = [row for case in cases for row in run_case(case, compressor, vector_store, reranker)]
    tiers = {str(row["tier"]) for row in results}
    required = {"short", "medium", "long"}
    if tiers != required:
        parser.error(
            "Benchmark cases must cover the 4–10 page, 20–30 page, and 50+ page tiers; "
            f"found {', '.join(sorted(tiers)) or 'none'}."
        )
    write_rows(results, args.workbook)
    print(json.dumps({"cases": len(cases), "rows": len(results), "tiers": sorted(tiers), "workbook": str(args.workbook)}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
