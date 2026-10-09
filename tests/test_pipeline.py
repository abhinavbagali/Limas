#!/usr/bin/env python3
"""Automated tests and formatted report generation for the context compressor project.

This script runs a focused test suite over the implemented compression pipeline and
writes a human-readable report to `reports/pipeline_test_results.txt`.
"""

from __future__ import annotations

import io
import asyncio
import os
import sys
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from context_compressor.compression.compressor import ContextCompressor
from context_compressor.compression.sentence_splitter import split_document_into_sentences
from context_compressor.coreference import DiscoursePreserver
from context_compressor.evaluation.benchmark import benchmark_query_set, benchmark_rag_pipeline
from context_compressor.evaluation.comparison_metrics import HEADERS, append_comparison_metrics
from context_compressor.gating import QueryComplexityGater
from context_compressor.models.cross_encoder import CrossEncoderRelevanceScorer
from context_compressor.models import cross_encoder as cross_encoder_module
from context_compressor.scorer import MultiHopInterdependenceScorer
from context_compressor.service.api import app
from context_compressor.service.groq_client import (
    GenerationResult,
    GroqConfigurationError,
    GroqRequestError,
)
from context_compressor.service.groq_client import generate_answer

RESULT_PATH = ROOT / "reports" / "pipeline_test_results.txt"


class DeterministicPairScorer:
    def __init__(self, scores: dict[str, float]):
        self.scores = scores

    def score_pairs(self, query: str, candidate_texts: list[str]) -> list[float]:
        return [self.scores.get(text, 0.0) for text in candidate_texts]


class LegacySingleScorer:
    def score_single(self, query: str, candidate: str) -> float:
        return 0.75 if "relevant" in candidate else 0.10


class ContextCompressorPipelineTests(unittest.TestCase):
    def test_sentence_splitter_basic(self):
        text = "Climate change is real. It affects rainfall. The planet warms."
        sentences = split_document_into_sentences(text)
        self.assertEqual(len(sentences), 3)
        self.assertTrue(all(sentence for sentence in sentences))
        self.assertEqual(sentences[0], "Climate change is real.")

    def test_relevance_scoring_prefers_relevant_text(self):
        with patch.object(CrossEncoderRelevanceScorer, "_load_model"):
            scorer = CrossEncoderRelevanceScorer()
        compressor = ContextCompressor(scorer=scorer)
        query = "What causes climate change?"
        relevant = "Greenhouse gases trap heat in the atmosphere and amplify warming."
        irrelevant = "The meeting starts at 4pm and the room is on the third floor."

        relevant_score = compressor.scorer.score_single(query, relevant)
        irrelevant_score = compressor.scorer.score_single(query, irrelevant)

        self.assertGreater(relevant_score, irrelevant_score)
        self.assertGreater(relevant_score, 0.0)

    def test_compress_keeps_relevant_sentences(self):
        compressor = ContextCompressor()
        query = "What causes climate change?"
        docs = [
            "Greenhouse gases trap heat in the atmosphere.",
            "The meeting is at 4pm.",
            "Human emissions increase the greenhouse effect.",
        ]

        result = compressor.compress(query=query, documents=docs, threshold=0.20, max_tokens=80)

        self.assertTrue(result.compressed_text)
        self.assertTrue(any("greenhouse" in sentence.lower() for sentence in result.selected_sentences))
        self.assertGreaterEqual(result.estimated_word_reduction_percent, 0.0)

    def test_sentence_audit_covers_retained_and_discarded_candidates(self):
        scorer = DeterministicPairScorer(
            {
                "Greenhouse gases trap heat.": 0.9,
                "The cafeteria closes at noon.": 0.05,
            }
        )
        result = ContextCompressor(
            scorer=scorer,
            discourse_preserver=DiscoursePreserver(use_spacy=False),
        ).compress(
            query="What causes greenhouse warming?",
            documents=[
                "Greenhouse gases trap heat.",
                "The cafeteria closes at noon.",
            ],
            threshold=0.5,
            enable_coref=False,
            enable_multihop=False,
            adaptive_gating=False,
        )

        self.assertEqual(len(result.sentence_audit), 2)
        retained, discarded = result.sentence_audit
        self.assertEqual(retained.sentence_index, 0)
        self.assertEqual(retained.isolated_score, 0.9)
        self.assertEqual(retained.pairwise_gain, 0.0)
        self.assertTrue(retained.retained)
        self.assertIsNotNone(retained.atomic_group_id)
        self.assertIn("threshold", retained.retention_reason)
        self.assertFalse(discarded.retained)
        self.assertIsNone(discarded.atomic_group_id)
        self.assertIn("Below threshold", discarded.discard_reason)

    def test_sentence_audit_assigns_one_group_to_preserved_antecedent(self):
        result = ContextCompressor(
            scorer=DeterministicPairScorer(
                {"Acme released its quarterly report.": 0.1, "It grew by 15%.": 0.9}
            ),
            discourse_preserver=DiscoursePreserver(use_spacy=False),
        ).compress(
            query="How much did Acme grow?",
            documents=["Acme released its quarterly report. It grew by 15%."],
            threshold=0.5,
            top_k=1,
            enable_multihop=False,
            adaptive_gating=False,
        )

        self.assertEqual([item.retained for item in result.sentence_audit], [True, True])
        self.assertEqual(
            result.sentence_audit[0].atomic_group_id,
            result.sentence_audit[1].atomic_group_id,
        )
        self.assertIn("antecedent", result.sentence_audit[0].retention_reason)

    def test_benchmark_returns_expected_shape(self):
        query = "How does climate change affect precipitation?"
        documents = [
            "Climate change alters the hydrological cycle by increasing evaporation.",
            "The office kitchen is closed on Fridays.",
            "Higher temperatures can shift storm tracks and modify rainfall patterns.",
        ]

        metrics = benchmark_rag_pipeline(
            query=query,
            documents=documents,
            answer="Climate change changes rainfall patterns by increasing evaporation and shifting storms.",
            threshold=0.25,
            top_k=4,
            max_tokens=80,
        )

        self.assertGreater(metrics.estimated_original_word_count, 0)
        self.assertGreaterEqual(metrics.estimated_compressed_word_count, 0)
        self.assertGreaterEqual(metrics.estimated_word_reduction_percent, 0.0)
        self.assertGreaterEqual(metrics.latency_ms, 0.0)
        self.assertGreaterEqual(metrics.baseline_latency_ms, 0.0)
        self.assertAlmostEqual(
            metrics.latency_overhead_ms,
            metrics.latency_ms - metrics.baseline_latency_ms,
            places=3,
        )

    def test_query_complexity_adapts_threshold_and_budget(self):
        gater = QueryComplexityGater()
        factual = gater.analyze("What is the capital of France?")
        complex_query = gater.analyze(
            "Compare Acme Corp with Beta Ltd: why did one acquire the other, "
            "and how did that change revenue?"
        )

        self.assertGreater(complex_query.complexity_score, factual.complexity_score)
        self.assertLess(complex_query.dynamic_threshold, factual.dynamic_threshold)
        self.assertGreater(complex_query.token_ratio_limit, factual.token_ratio_limit)
        self.assertEqual(complex_query.intent_type, "comparative")
        self.assertTrue(0.30 <= complex_query.dynamic_threshold <= 0.75)
        self.assertTrue(0.20 <= complex_query.token_ratio_limit <= 0.80)

    def test_discourse_analyzer_links_pronouns_and_connectives_within_passage(self):
        analyzer = DiscoursePreserver(use_spacy=False)
        sentences = [
            "Acme released its quarterly earnings report.",
            "It grew by 15% in Q3.",
            "However, investors expected stronger revenue.",
            "The water is stored in a reservoir.",
            "They use it for irrigation.",
        ]
        links = analyzer.find_links(
            sentences,
            document_ids=[0, 0, 0, 1, 1],
        )

        self.assertEqual(
            [(link.antecedent_index, link.dependent_index) for link in links],
            [(0, 1), (1, 2), (3, 4)],
        )
        retained = analyzer.mandatory_antecedents({2}, links)
        self.assertEqual(retained, {0, 1, 2})

    def test_pairwise_scorer_promotes_conditional_followup(self):
        scorer = DeterministicPairScorer(
            {
                "Cause sentence.": 0.10,
                "Conditional result.": 0.20,
                "Cause sentence. Conditional result.": 0.90,
            }
        )
        result = MultiHopInterdependenceScorer(scorer).score(
            "What happened?",
            ["Cause sentence.", "Conditional result."],
            [0, 0],
        )

        self.assertGreater(result.weights[1], 0.20)
        self.assertEqual(result.pair_dependencies, [(0, 1)])
        self.assertEqual(len(result.weights), 2)

    def test_pairwise_scorer_supports_legacy_single_only_scorers(self):
        result = MultiHopInterdependenceScorer(LegacySingleScorer()).score(
            "query",
            ["relevant evidence.", "unrelated sentence."],
            [0, 0],
            enabled=False,
        )
        self.assertEqual(result.weights, [0.75, 0.10])

    def test_single_logit_cross_encoder_scores_are_monotonic(self):
        torch = cross_encoder_module.torch
        if torch is None:
            self.skipTest("PyTorch is unavailable.")
        with patch.object(CrossEncoderRelevanceScorer, "_load_model"):
            scorer = CrossEncoderRelevanceScorer(batch_size=2)

        class Tokenizer:
            def __call__(self, queries, candidates, **kwargs):
                return {"input_ids": torch.tensor([[0], [1]])}

        class Model:
            def __call__(self, input_ids):
                return SimpleNamespace(logits=torch.tensor([[-2.0], [2.0]]))

        scorer.tokenizer = Tokenizer()
        scorer.model = Model()
        scorer.fallback_mode = False
        scores = scorer.score_pairs("query", ["low", "high"])
        self.assertLess(scores[0], scores[1])
        self.assertAlmostEqual(scores[0], 1.0 / (1.0 + 2.718281828 ** 2), places=5)

    def test_compressor_keeps_atomic_coreference_pair_and_reports_metadata(self):
        scorer = DeterministicPairScorer(
            {
                "Acme released earnings.": 0.05,
                "It grew by 15% in Q3.": 0.20,
                "Acme released earnings. It grew by 15% in Q3.": 0.90,
            }
        )
        compressor = ContextCompressor(
            scorer=scorer,
            discourse_preserver=DiscoursePreserver(use_spacy=False),
        )
        result = compressor.compress(
            query="Acme revenue?",
            documents=["Acme released earnings. It grew by 15% in Q3."],
            threshold=0.30,
            top_k=1,
            max_tokens=100,
        )

        self.assertEqual(result.retained_sentence_indices, [0, 1])
        self.assertEqual(result.preserved_coreferences_count, 1)
        self.assertIn("Acme released earnings.", result.compressed_text)
        self.assertIn("It grew by 15%", result.compressed_text)
        self.assertTrue(0.0 <= result.query_complexity_score <= 1.0)
        self.assertTrue(0.30 <= result.dynamic_threshold_used <= 0.75)
        self.assertTrue(0.20 <= result.dynamic_token_ratio_limit <= 0.80)
        self.assertGreater(result.pairwise_scoring_latency_ms, 0.0)

    def test_benchmark_query_set_groups_latency_and_quality_proxies(self):
        compressor = ContextCompressor(
            scorer=DeterministicPairScorer(
                {
                    "France is the capital of France.": 0.9,
                    "Acme bought Beta.": 0.7,
                    "It increased revenue.": 0.7,
                    "Acme bought Beta. It increased revenue.": 0.9,
                }
            ),
            discourse_preserver=DiscoursePreserver(use_spacy=False),
        )
        report = benchmark_query_set(
            [
                {
                    "query": "What is the capital of France?",
                    "documents": ["France is the capital of France."],
                    "answer": "France is the capital of France.",
                },
                {
                    "query": "Compare Acme and Beta: how did the acquisition affect revenue?",
                    "documents": ["Acme bought Beta. It increased revenue."],
                    "answer": "Acme bought Beta and revenue increased.",
                },
            ],
            compressor=compressor,
        )

        self.assertEqual(len(report["cases"]), 2)
        self.assertEqual(report["summary_by_intent"]["factual"]["case_count"], 1)
        self.assertEqual(report["summary_by_intent"]["comparative"]["case_count"], 1)
        self.assertIn("mean_latency_overhead_ms", report["summary_by_intent"]["factual"])
        self.assertIn("mean_answer_similarity_gain", report["summary_by_intent"]["comparative"])

    def test_api_health_and_compression(self):
        client = TestClient(app)

        ui_response = client.get("/")
        self.assertEqual(ui_response.status_code, 200)
        self.assertIn("Baseline RAG · Without compressor", ui_response.text)
        self.assertIn("Compressed RAG · With compressor", ui_response.text)
        self.assertIn('fetch("/compare"', ui_response.text)
        self.assertIn("GROQ_API_KEY", ui_response.text)
        self.assertIn("theme-toggle", ui_response.text)
        self.assertIn("<title>Limas — Subcortical Context Compressor</title>", ui_response.text)
        self.assertIn("Subcortical Context Compression", ui_response.text)
        self.assertIn('id="upload-dropzone"', ui_response.text)
        self.assertIn('fetch("/analytics")', ui_response.text)
        self.assertIn('id="test-case"', ui_response.text)
        self.assertIn('id="heatmap"', ui_response.text)
        self.assertIn("enable-coref", ui_response.text)
        self.assertIn("adaptive-gating", ui_response.text)

        health_response = client.get("/health")
        self.assertEqual(health_response.status_code, 200)
        self.assertEqual(health_response.json()["status"], "ok")
        with patch.dict(os.environ, {"GROQ_API_KEY": "test-key"}):
            config_response = client.get("/config")
        self.assertEqual(config_response.status_code, 200)
        self.assertEqual(
            config_response.json(),
            {"groq_configured": True, "model": "openai/gpt-oss-120b"},
        )
        self.assertNotIn("test-key", config_response.text)
        self.assertEqual(client.get("/health").json()["service"], "limas")

        payload = {
            "query": "What causes climate change?",
            "documents": [
                "Greenhouse gases trap heat in the atmosphere.",
                "The meeting is at 4pm.",
                "Human emissions increase the greenhouse effect.",
            ],
            "threshold": 0.2,
            "max_tokens": 120,
        }

        compress_response = client.post("/compress", json=payload)
        self.assertEqual(compress_response.status_code, 200)
        body = compress_response.json()
        self.assertIn("compressed_text", body)
        self.assertIn("selected_sentences", body)
        self.assertIn("estimated_word_reduction_percent", body)
        self.assertIn("query_complexity_score", body)
        self.assertIn("dynamic_threshold_used", body)
        self.assertIn("preserved_coreferences_count", body)
        self.assertIn("retained_sentence_indices", body)
        self.assertEqual(len(body["sentence_audit"]), 3)
        self.assertEqual(
            {item["sentence_index"] for item in body["sentence_audit"]},
            {0, 1, 2},
        )
        self.assertEqual(body["dynamic_threshold_used"], 0.2)
        self.assertTrue(0.20 <= body["dynamic_token_ratio_limit"] <= 0.80)

        disabled_response = client.post(
            "/compress",
            json={
                "query": "What is this?",
                "documents": ["This is a finding. It matters."],
                "enable_coref": False,
                "enable_multihop": False,
                "adaptive_gating": False,
            },
        )
        self.assertEqual(disabled_response.status_code, 200, disabled_response.text)
        disabled_body = disabled_response.json()
        self.assertEqual(disabled_body["preserved_coreferences_count"], 0)
        self.assertEqual(disabled_body["pairwise_scoring_latency_ms"], 0.0)
        self.assertIsNone(disabled_body["dynamic_token_ratio_limit"])

    def test_upload_endpoint_extracts_text_files_and_pdf_pages(self):
        client = TestClient(app)
        text_response = client.post(
            "/upload_pdf",
            files={"file": ("sources.txt", b"First passage.\n\nSecond passage.", "text/plain")},
        )
        self.assertEqual(text_response.status_code, 200, text_response.text)
        self.assertEqual(
            text_response.json()["passages"],
            ["First passage.", "Second passage."],
        )
        self.assertEqual(text_response.json()["page_count"], 1)

        fake_reader = SimpleNamespace(
            pages=[
                SimpleNamespace(extract_text=lambda: "First PDF page."),
                SimpleNamespace(extract_text=lambda: "Second PDF page."),
            ]
        )
        with patch("context_compressor.service.api.PdfReader", return_value=fake_reader):
            pdf_response = client.post(
                "/upload_pdf",
                files={"file": ("sources.PDF", b"%PDF-test", "application/pdf")},
            )
        self.assertEqual(pdf_response.status_code, 200, pdf_response.text)
        self.assertEqual(
            pdf_response.json()["passages"],
            ["First PDF page.", "Second PDF page."],
        )
        self.assertEqual(pdf_response.json()["page_count"], 2)

        unsupported = client.post(
            "/upload_pdf",
            files={"file": ("sources.docx", b"content", "application/octet-stream")},
        )
        self.assertEqual(unsupported.status_code, 415)

    def test_analytics_endpoint_aggregates_saved_workbook_metrics(self):
        client = TestClient(app)
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "metrics.xlsx"
            metric_row = {
                "Test Case": "climate",
                "Without Compressor - Prompt Tokens (Groq actual)": 100,
                "With Compressor - Prompt Tokens (Groq actual)": 60,
                "Prompt Tokens Saved (Groq actual)": 40,
                "Prompt Token Reduction (%) (Groq actual)": 40.0,
                "Latency Saved (ms)": 50.0,
                "Without Compressor - Total Latency (ms)": 100.0,
                "With Compressor - Reference Similarity": 0.9,
                "Preserved Coreferences": 2,
                "Query Intent Type": "comparative",
                "Dynamic Threshold Used": 0.4,
            }
            with patch("context_compressor.service.api.DEFAULT_WORKBOOK", workbook_path):
                append_comparison_metrics(metric_row, workbook_path=workbook_path)
                response = client.get("/analytics")

            self.assertEqual(response.status_code, 200, response.text)
            body = response.json()
            self.assertEqual(body["runs_total"], 1)
            self.assertEqual(body["avg_prompt_token_reduction_percent"], 40.0)
            self.assertEqual(body["avg_latency_saved_ms"], 50.0)
            self.assertEqual(body["avg_quality_retention_percent"], 90.0)
            self.assertEqual(body["coreferences_preserved"], 2)
            self.assertEqual(body["intent_distribution"], [{"name": "Comparative", "count": 1}])
            self.assertEqual(body["activity"][0]["saved_percent"], 40.0)
            self.assertEqual(body["performance"][0]["baseline_prompt_tokens"], 100)
            self.assertEqual(body["dynamic_threshold_range"], {"min": 0.4, "max": 0.4})

    def test_compare_generates_both_modes_and_saves_metrics(self):
        client = TestClient(app)
        generated = AsyncMock(
            side_effect=[
                GenerationResult("Baseline answer.", 120.0, 100, 12),
                GenerationResult("Compressed answer.", 80.0, 55, 10),
            ]
        )
        payload = {
            "query": "What causes climate change?",
            "documents": [
                "Greenhouse gases trap heat in the atmosphere.",
                "The meeting is at 4pm.",
                "Human emissions increase the greenhouse effect.",
            ],
            "expected_answer": "Greenhouse gas emissions trap heat.",
            "test_case": "climate smoke test",
            "threshold": 0.2,
            "max_tokens": 120,
        }

        with patch("context_compressor.service.api.generate_answer", generated), patch(
            "context_compressor.service.api.append_comparison_metrics"
        ) as save_metrics:
            response = client.post("/compare", json=payload)

        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertEqual(body["without_compressor_answer"], "Baseline answer.")
        self.assertEqual(body["with_compressor_answer"], "Compressed answer.")
        self.assertEqual(body["without_compressor_prompt_tokens"], 100)
        self.assertEqual(body["with_compressor_prompt_tokens"], 55)
        self.assertEqual(body["actual_prompt_tokens_saved"], 45)
        self.assertAlmostEqual(body["actual_prompt_token_reduction_percent"], 45.0)
        self.assertTrue(body["metrics_saved"])
        self.assertEqual(generated.await_count, 2)
        self.assertNotEqual(generated.await_args_list[0].args[1], generated.await_args_list[1].args[1])
        metric_values = save_metrics.call_args.args[0]
        self.assertIn("Without Compressor - LLM Latency (ms)", metric_values)
        self.assertIn("With Compressor - Prompt Tokens (Groq actual)", metric_values)
        self.assertEqual(metric_values["Prompt Tokens Saved (Groq actual)"], 45)
        self.assertAlmostEqual(metric_values["Prompt Token Reduction (%) (Groq actual)"], 45.0)
        self.assertEqual(len(body["sentence_audit"]), 3)

    def test_compare_runs_generation_requests_concurrently(self):
        client = TestClient(app)
        active = 0
        max_active = 0

        async def generate(query, context):
            nonlocal active, max_active
            active += 1
            max_active = max(max_active, active)
            await asyncio.sleep(0.02)
            active -= 1
            return GenerationResult("Grounded answer.", 20.0, 10, 4)

        with patch("context_compressor.service.api.generate_answer", generate), patch(
            "context_compressor.service.api.append_comparison_metrics"
        ):
            response = client.post(
                "/compare",
                json={
                    "query": "Question?",
                    "documents": ["First source sentence.", "Second source sentence."],
                },
            )

        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(max_active, 2)

    def test_compare_reports_branch_for_failed_concurrent_request(self):
        client = TestClient(app)
        generated = AsyncMock(
            side_effect=[
                GroqRequestError("baseline provider failure"),
                GenerationResult("Compressed answer.", 15.0, 8, 3),
            ]
        )
        with patch("context_compressor.service.api.generate_answer", generated):
            response = client.post(
                "/compare",
                json={"query": "Question?", "documents": ["Context sentence."]},
            )

        self.assertEqual(response.status_code, 502)
        self.assertIn("baseline", response.json()["detail"])
        self.assertIn("baseline provider failure", response.json()["detail"])

    def test_compare_reports_missing_groq_key(self):
        client = TestClient(app)
        with patch(
            "context_compressor.service.api.generate_answer",
            new_callable=AsyncMock,
            side_effect=GroqConfigurationError("GROQ_API_KEY is not configured."),
        ):
            response = client.post(
                "/compare",
                json={"query": "Question?", "documents": ["Context sentence."]},
            )
        self.assertEqual(response.status_code, 503)
        self.assertIn("GROQ_API_KEY", response.json()["detail"])

    def test_missing_groq_usage_is_not_replaced_with_word_estimate(self):
        client = TestClient(app)
        generated = AsyncMock(
            side_effect=[
                GenerationResult("Baseline.", 20.0, None, None),
                GenerationResult("Compressed.", 15.0, None, None),
            ]
        )
        with patch("context_compressor.service.api.generate_answer", generated), patch(
            "context_compressor.service.api.append_comparison_metrics"
        ):
            response = client.post(
                "/compare",
                json={"query": "Question?", "documents": ["A context sentence."]},
            )
        self.assertEqual(response.status_code, 200, response.text)
        body = response.json()
        self.assertIsNone(body["without_compressor_prompt_tokens"])
        self.assertIsNone(body["with_compressor_prompt_tokens"])
        self.assertIsNone(body["actual_prompt_tokens_saved"])
        self.assertIsNone(body["actual_prompt_token_reduction_percent"])
        self.assertIn("estimated_original_word_count", body)

    def test_groq_client_uses_server_key_and_expected_model(self):
        response = MagicMock()
        response.is_success = True
        response.json.return_value = {
            "choices": [{"message": {"content": "Grounded answer."}}],
            "usage": {"prompt_tokens": 80, "completion_tokens": 9},
        }
        client = AsyncMock()
        client.post = AsyncMock(return_value=response)
        client_manager = MagicMock()
        client_manager.__aenter__ = AsyncMock(return_value=client)
        client_manager.__aexit__ = AsyncMock(return_value=False)

        with patch.dict(
            os.environ,
            {"GROQ_API_KEY": "test-secret", "GROQ_MODEL": "openai/gpt-oss-120b"},
        ), patch(
            "context_compressor.service.groq_client.httpx.AsyncClient",
            return_value=client_manager,
        ):
            result = asyncio.run(generate_answer("Question?", "Relevant context."))

        self.assertEqual(result.answer, "Grounded answer.")
        self.assertEqual(result.prompt_tokens, 80)
        self.assertEqual(result.completion_tokens, 9)
        call_kwargs = client.post.call_args.kwargs
        self.assertEqual(call_kwargs["headers"]["Authorization"], "Bearer test-secret")
        self.assertEqual(call_kwargs["json"]["model"], "openai/gpt-oss-120b")

    def test_comparison_workbook_migrates_legacy_metrics(self):
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "eval_met.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "eval_met"
            worksheet.append(
                [
                    "Timestamp (UTC)",
                    "Query",
                    "Original Tokens",
                    "Compressed Tokens",
                    "Token Reduction (%)",
                    "Compression Latency (ms)",
                    "Answer Similarity",
                ]
            )
            worksheet.append(["2026-01-01T00:00:00Z", "Legacy query", 20, 10, 50.0, 4.0, 0.5])
            workbook.save(workbook_path)

            append_comparison_metrics(
                {
                    "Query": "New query",
                    "Model": "openai/gpt-oss-120b",
                    "Query Complexity Score": 0.7,
                    "Query Intent Type": "multi_hop",
                    "Retained Sentence Indices": "1,2",
                },
                workbook_path,
            )
            saved = load_workbook(workbook_path, read_only=True, data_only=True)
            rows = list(saved["eval_met"].values)
            saved.close()

            self.assertEqual(rows[0], HEADERS)
            self.assertEqual(rows[1][2], "Legacy query")
            self.assertEqual(rows[1][23:27], (20, 10, 50.0, 0.5))
            self.assertEqual(rows[2][2], "New query")
            self.assertEqual(rows[2][3], "openai/gpt-oss-120b")
            self.assertEqual(len(rows[0]), 37)
            self.assertEqual(rows[2][27:30], (0.7, "multi_hop", None))
            self.assertEqual(rows[2][33], "1,2")

    def test_comparison_workbook_migrates_previous_comparison_schema(self):
        with tempfile.TemporaryDirectory() as directory:
            workbook_path = Path(directory) / "eval_met.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "eval_met"
            worksheet.append(
                [
                    "Timestamp (UTC)",
                    "Test Case",
                    "Query",
                    "Model",
                    "Retrieved Passages",
                    "Original Context Tokens (estimated)",
                    "Compressed Context Tokens (estimated)",
                    "Context Token Reduction (%)",
                    "Sentences Retained",
                    "Compression Latency (ms)",
                    "Without Compressor - LLM Latency (ms)",
                    "With Compressor - LLM Latency (ms)",
                    "Without Compressor - Total Latency (ms)",
                    "With Compressor - Total Latency (ms)",
                    "Latency Saved (ms)",
                    "Without Compressor - Prompt Tokens (API)",
                    "With Compressor - Prompt Tokens (API)",
                    "Without Compressor - Completion Tokens (API)",
                    "With Compressor - Completion Tokens (API)",
                    "Without Compressor - Answer",
                    "With Compressor - Answer",
                    "Expected Answer",
                    "Without Compressor - Reference Similarity",
                    "With Compressor - Reference Similarity",
                    "Legacy Answer Similarity (old benchmark)",
                ]
            )
            worksheet.append(
                [
                    "2026-01-01T00:00:00Z",
                    "old run",
                    "Old query",
                    "openai/gpt-oss-120b",
                    2,
                    40,
                    20,
                    50.0,
                    3,
                    12.0,
                    90.0,
                    70.0,
                    90.0,
                    82.0,
                    8.0,
                    120,
                    80,
                    10,
                    9,
                    "Baseline",
                    "Compressed",
                    "Expected",
                    0.5,
                    0.6,
                    None,
                ]
            )
            workbook.save(workbook_path)

            append_comparison_metrics({"Query": "New query"}, workbook_path)
            saved = load_workbook(workbook_path, read_only=True, data_only=True)
            rows = list(saved["eval_met"].values)
            saved.close()

            self.assertEqual(rows[0], HEADERS)
            self.assertEqual(rows[1][2], "Old query")
            self.assertEqual(rows[1][5:7], (120, 80))
            self.assertEqual(rows[1][7], 40)
            self.assertAlmostEqual(rows[1][8], 100.0 / 3.0)
            self.assertEqual(rows[1][23:26], (40, 20, 50.0))
            self.assertEqual(rows[2][2], "New query")


def run_suite() -> tuple[int, str]:
    suite = unittest.defaultTestLoader.loadTestsFromModule(sys.modules[__name__])
    stream = io.StringIO()
    runner = unittest.TextTestRunner(stream=stream, verbosity=2)
    result = runner.run(suite)

    timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    summary = (
        "=== Limas Pipeline Test Report ===\n"
        f"Timestamp: {timestamp}\n\n"
        f"{stream.getvalue().strip()}\n\n"
        f"Summary: {result.testsRun} tests run, {len(result.failures)} failed, {len(result.errors)} errors, {len(result.skipped)} skipped\n"
    )

    if not result.wasSuccessful():
        summary += "Status: FAILED\n"
    else:
        summary += "Status: PASS\n"

    RESULT_PATH.parent.mkdir(parents=True, exist_ok=True)
    RESULT_PATH.write_text(summary, encoding="utf-8")
    return (0 if result.wasSuccessful() else 1), summary


if __name__ == "__main__":
    exit_code, summary = run_suite()
    print(summary)
    raise SystemExit(exit_code)
